#!/usr/bin/env python3
# build_doc.py - part of the Sophos Mobile to Wazuh integration
# Copyright (C) 2026 Gerardo Caceres
#
# This program is free software; you can redistribute it and/or modify it
# under the terms of the GNU General Public License version 2 as published
# by the Free Software Foundation.
#
# This program is distributed in the hope that it will be useful, but
# WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the GNU General
# Public License for more details.
#
# You should have received a copy of the GNU General Public License along
# with this program; if not, see <https://www.gnu.org/licenses/>.
#
# Renders the customer-facing technical document as a .docx.
#
#   pip install python-docx
#   python3 docs/build_doc.py "Sophos Mobile integration with Wazuh.docx"
#
# The prose lives in doc_content.py. The appendices read the real source files
# out of the repository at build time, so the document cannot drift from the
# code it documents.
#
# Visual style follows the existing Wazuh document template: Manrope, US
# Letter with 1in margins, numbered H1/H2/H3 in #0B0B0C, Wazuh-blue table
# headers, #F6F6F6 banding and #F2F7FE callouts.
import os
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

OUT = sys.argv[1] if len(sys.argv) > 1 else "Sophos Mobile integration with Wazuh.docx"

BLUE = "3585F9"
GREY = "F6F6F6"
LIGHTBLUE = "F2F7FE"
INK = "0B0B0C"
BODY = "1B1D1F"
CODEC = "24292F"
LINE = "E3E5E8"
MONO = "Consolas"
FONT = "Manrope"
CW = 9360  # content width in twips: 12240 - 1440*2

doc = Document()

# ----------------------------------------------------------------- page setup
sec = doc.sections[0]
sec.page_width = Twips(12240)
sec.page_height = Twips(15840)
for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, attr, Twips(1440))

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BODY)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.15

for name, pt in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(pt)
    st.font.bold = True
    st.font.italic = False
    st.font.color.rgb = RGBColor.from_string(INK)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st.paragraph_format.space_before = Pt(14 if pt == 17 else 11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number"):
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(BODY)
    st.paragraph_format.space_after = Pt(4)


# --------------------------------------------------------------- xml helpers
def _shd(fill):
    e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear")
    e.set(qn("w:color"), "auto")
    e.set(qn("w:fill"), fill)
    return e


def cell_shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(_shd(fill))


def par_shade(par, fill):
    par._p.get_or_add_pPr().append(_shd(fill))


def par_left_border(par, color, size=18):
    bd = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    bd.append(left)
    par._p.get_or_add_pPr().append(bd)


def table_borders(tbl, color=LINE, size=2):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tbl._tbl.tblPr.append(borders)


def cell_borders_left_only(cell, color, size=24):
    bd = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "nil")
        bd.append(e)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:color"), color)
    bd.append(left)
    cell._tc.get_or_add_tcPr().append(bd)


def cell_margins(cell, top=90, bottom=90, left=135, right=135):
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom),
                     ("left", left), ("right", right)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)


def repeat_header(row):
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(h)


# Ask Word to refresh the table of contents when the document is opened.
# python-docx does not expose this, and Word strips it again once it has
# updated the fields, which is why it is set explicitly here.
_uf = OxmlElement("w:updateFields")
_uf.set(qn("w:val"), "true")
doc.settings.element.append(_uf)


# ------------------------------------------------------------------ rendering
def norm(parts):
    """Accept a string, or a list of (text, flags) with flags in b/i/c."""
    if isinstance(parts, str):
        return [(parts, "")]
    out = []
    for item in parts:
        if isinstance(item, str):
            out.append((item, ""))
        else:
            out.append((item[0], item[1] if len(item) > 1 else ""))
    return out


def emit(par, parts):
    for text, flags in norm(parts):
        r = par.add_run(text)
        if "b" in flags:
            r.bold = True
        if "i" in flags:
            r.italic = True
        if "c" in flags:
            r.font.name = MONO
            r._element.rPr.rFonts.set(qn("w:hAnsi"), MONO)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(CODEC)
    return par


def op_title(text):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(24)
    par.paragraph_format.space_after = Pt(6)
    r = par.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)


def op_subtitle(text):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(22)
    r = par.add_run(text)
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string("5A6169")


def op_code(lines, size=8.5):
    for i, ln in enumerate(lines):
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_before = Pt(4 if i == 0 else 0)
        pf.space_after = Pt(8 if i == len(lines) - 1 else 0)
        pf.left_indent = Twips(140)
        pf.right_indent = Twips(140)
        pf.line_spacing = 1.0
        pf.widow_control = False
        par_shade(par, GREY)
        par_left_border(par, "C9CDD2")
        r = par.add_run(ln if ln else " ")
        r.font.name = MONO
        r._element.rPr.rFonts.set(qn("w:hAnsi"), MONO)
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(CODEC)


def op_callout(lead, bodytext, fill=LIGHTBLUE):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Twips(CW)
    cell_shade(cell, fill)
    cell_borders_left_only(cell, BLUE)
    cell_margins(cell, 150, 150, 200, 200)
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    r = par.add_run(lead + " ")
    r.bold = True
    emit(par, bodytext)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def op_table(headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    table_borders(tbl)
    hdr = tbl.rows[0]
    repeat_header(hdr)
    for i, htext in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Twips(widths[i])
        cell_shade(cell, BLUE)
        cell_margins(cell)
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        r = par.add_run(htext)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string("FFFFFF")
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, parts in enumerate(row):
            cell = cells[ci]
            cell.width = Twips(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            if ri % 2 == 1:
                cell_shade(cell, GREY)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            emit(par, parts)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_toc():
    par = doc.add_paragraph()
    run = par.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = r'TOC \o "1-3" \h \z \u'
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Update this field in Word to build the table of contents."
    f2.append(t)
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, f3):
        run._r.append(e)


def add_hyperlink(par, text, url):
    r_id = par.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "1A6DD8")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(col)
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    hl.append(r)
    par._p.append(hl)


def render(ops):
    for op in ops:
        kind = op[0]
        if kind == "title":
            op_title(op[1])
        elif kind == "sub":
            op_subtitle(op[1])
        elif kind in ("h1", "h2", "h3"):
            doc.add_heading(op[1], level=int(kind[1]))
        elif kind == "p":
            emit(doc.add_paragraph(), op[1])
        elif kind == "bullet":
            emit(doc.add_paragraph(style="List Bullet"), op[1])
        elif kind == "step":
            emit(doc.add_paragraph(style="List Number"), op[1])
        elif kind == "code":
            op_code(op[1])
        elif kind == "srccode":
            op_code(op[1], op[2] if len(op) > 2 else 7.5)
        elif kind == "callout":
            op_callout(op[1], op[2])
        elif kind == "table":
            op_table(op[1], op[2], op[3])
        elif kind == "toc":
            add_toc()
        elif kind == "break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "spacer":
            doc.add_paragraph().paragraph_format.space_after = Pt(op[1])
        elif kind == "link":
            add_hyperlink(doc.add_paragraph(style="List Bullet"), op[1], op[2])
        else:
            raise ValueError("unknown op " + kind)


from doc_content import OPS  # noqa: E402

render(OPS)
doc.save(OUT)
print("written:", OUT)
